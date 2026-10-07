import csv
import os
import re
import sqlite3
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
import zipfile

from database import (
    delete_product_backlog,
    get_fibonacci,
    get_sprints,
    get_status_product_backlog,
    insert_product_backlog,
    list_product_backlog,
    reset_project_data,
    update_product_backlog,
)
from grid_sizing import fit_tree_columns


MODULE_LABEL = "Módulo"
PANEL_STYLE = "Panel.TFrame"
MUTED_LABEL_STYLE = "Muted.TLabel"
SECTION_LABEL_STYLE = "Section.TLabel"
PRIMARY_BUTTON_STYLE = "Primary.TButton"
NEW_STORY_BUTTON_STYLE = "NewStory.TButton"
ICON_BUTTON_STYLE = "Icon.TButton"
BUTTON_WIDTH = 18
ICON_BUTTON_WIDTH = 4
FONT_NAME = "Segoe UI"
DEPENDENCY_ERROR_TITLE = "Dependência ausente"
DEPENDENCY_ERROR_MESSAGE = "Instale as dependências com: pip install -r requirements.txt"
EXPORT_COLUMNS = ["Id", MODULE_LABEL, "História", "Prioridade", "Story Points", "Sprint", "Status", "Observações"]
BACKLOG_GRID_COLUMNS = (
    ("id", "Id", 0, 90), ("module", MODULE_LABEL, 1, 130),
    ("story", "História de usuário", 2, 220), ("priority", "Prioridade", 3, 110),
    ("points", "Pontos", 4, 75), ("sprint", "Sprint", 5, 110), ("status", "Status", 6, 130),
)
FORM_SUMMARIES = {
    "Dashboards": "Visualize cards e gráficos consolidados dos seis formulários para decisões.",
    "Product Backlog": "Organize histórias, módulos, prioridade, pontos e status do produto.",
    "Sprints Backlog": "Divida histórias em tarefas e acompanhe responsáveis, horas e status.",
    "Histórico de Sprints": "Registre pontos planejados/entregues, velocidade e capacidade por tarefa.",
    "Governança": "Cadastre políticas, responsáveis, periodicidade e histórias relacionadas.",
    "Transparência": "Registre relatórios, frequência, responsável, visibilidade e histórias relacionadas.",
    "Qualidade": "Acompanhe metas, valores reais, atingimento e histórias relacionadas.",
}
FORM_MANUALS = {
    "Dashboards": "OBJETIVO\nConsolidar os dados dos seis formulários em indicadores visuais para tomada de decisão.\n\nABAS\nCada aba representa um formulário: Product Backlog, Sprints Backlog, Histórico, Governança, Transparência e Qualidade.\n\nCARDS\nMostram totais e percentuais-chave: histórias, pontos, tarefas, horas, velocity média, políticas ativas, relatórios, índice de qualidade e histórias vinculadas.\n\nGRÁFICOS\nDistribuição por status, pontos por prioridade, horas por sprint, velocity por sprint, planejado vs entregue, políticas por periodicidade, relatórios por visibilidade e atingimento por métrica.\n\nAÇÕES\nUse Atualizar para recarregar os dados do banco após alterações nos formulários.",
    "Product Backlog": "OBJETIVO\nOrganizar as necessidades do produto em histórias priorizadas.\n\nDADOS\nId, módulo, história de usuário, prioridade, Story Points, sprint alocada, status e observações.\n\nINDICADORES\nStory Points segue a escala Fibonacci; a ajuda da prioridade mostra significado e exemplos. Status mostra significado e quando usar.\n\nAÇÕES\nNova história inicia cadastro; Salvar grava ou atualiza; Editar/Excluir atuam sobre a linha selecionada. Busca, status e sprint filtram o grid. CSV/Excel importam e exportam os registros.",
    "Sprints Backlog": "OBJETIVO\nDetalhar as tarefas técnicas de cada história e acompanhar a execução da sprint.\n\nDADOS\nId_Task, título da sprint, User_Story, subtarefa, responsável, Story_Points, horas estimadas/gastas/restantes, status e data de conclusão.\n\nAUTOMAÇÕES\nUser_Story traz Id e módulo do Product Backlog; Story_Points vem da história. Horas restantes = máximo(0, estimadas - gastas), ou 00:00 quando concluída.\n\nAÇÕES\nSelecione ou pesquise tarefas para editar/excluir. Nova tarefa prepara um registro; Salvar persiste no SQLite. Os botões acima do grid importam e exportam CSV/Excel; linhas com Id_Task já existente são ignoradas.",
    "Histórico de Sprints": "OBJETIVO\nManter o histórico de planejamento e entrega das tarefas em cada sprint.\n\nDADOS\nId_Historico, Id_Task, sprint, User_Story, datas, pontos planejados/entregues, velocidade, capacidade e status.\n\nAUTOMAÇÕES\nId_Task carrega sprint, história, pontos e horas estimadas de Sprints_Backlog. Pontos entregues iguala os planejados quando Status_Sprint é CONCLUÍDO; nos demais casos é zero. Velocity_% = entregues / planejados × 100; se planejados for zero, 0,0%.\n\nAÇÕES\nUse a busca, filtros e botões do grid para localizar, editar ou excluir históricos. Os botões acima do grid importam e exportam CSV/Excel; o Id_Historico é gerado automaticamente na importação.",
    "Governança": "OBJETIVO\nRegistrar políticas, responsáveis, periodicidade e situação de governança.\n\nVÍNCULOS\nSelecione uma ou mais histórias do Product_Backlog para associar à política. Os vínculos são salvos em Backlog_Governanca.\n\nAÇÕES\nNova política inicia cadastro; Salvar grava ou atualiza; Editar/Excluir atuam sobre a linha selecionada. A busca filtra o grid. Os botões acima do grid importam e exportam CSV/Excel; políticas com nome já existente são ignoradas.",
    "Transparência": "OBJETIVO\nRegistrar os relatórios publicados, frequência, responsável e público de visibilidade.\n\nVÍNCULOS\nSelecione uma ou mais histórias do Product_Backlog; as referências ficam em Backlog_Transparencia e aparecem no grid.\n\nAÇÕES\nNovo relatório inicia cadastro; Salvar grava ou atualiza; Editar/Excluir atuam sobre a linha selecionada. A busca filtra o grid. Os botões acima do grid importam e exportam CSV/Excel; relatórios com nome já existente são ignorados.",
    "Qualidade": "OBJETIVO\nAcompanhar métricas de qualidade comparando Valor_Real com Valor_Alvo.\n\nINDICADORES\nAtingimento por métrica = mínimo(Valor_Real / Valor_Alvo, 1) × 100. O índice do projeto é a média dos atingimentos.\n\nFONTES\nVelocity e tempo médio podem vir de tarefas concluídas em Sprints_Backlog. Coverage aceita relatório Coverage.py JSON/XML. Disponibilidade aceita CSV online_hours,total_hours ou online_seconds,total_seconds. Sem fonte externa, Valor_Real é manual.\n\nVÍNCULOS\nAssocie histórias do Product_Backlog; referências ficam em Backlog_Qualidade.\n\nAÇÕES\nOs botões acima do grid importam e exportam CSV/Excel; métricas com nome já existente são ignoradas.",
}


class BacklogApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Scrum | Product Backlog")
        self.root.geometry("1380x820")
        self.root.minsize(1080, 680)
        self.root.configure(bg="#edf1ee")

        self.status_data = get_status_product_backlog()
        self.fibonacci_data = get_fibonacci()
        self.status_options = list(self.status_data)
        self.sprint_options = get_sprints()
        self.points_options = list(self.fibonacci_data) or ["1", "2", "3", "5", "8", "13", "21"]
        self.records = []
        self.editing_id = None
        self.search_mode = "all"

        self.search_var = tk.StringVar()
        self.status_filter_var = tk.StringVar(value="TODOS")
        self.sprint_filter_var = tk.StringVar(value="TODAS")
        self.id_var = tk.StringVar()
        self.module_var = tk.StringVar()
        self.priority_var = tk.StringVar()
        self.points_var = tk.StringVar()
        self.sprint_var = tk.StringVar()
        self.status_var = tk.StringVar()

        self._configure_style()
        self._build_layout()
        self.refresh_records()
        self.show_main_menu()

    def _configure_style(self):
        style = ttk.Style(self.root)
        style.theme_use("clam")
        style.configure("TFrame", background="#edf1ee")
        style.configure(PANEL_STYLE, background="#ffffff")
        style.configure("TLabel", background="#ffffff", foreground="#23352f", font=(FONT_NAME, 9))
        style.configure(MUTED_LABEL_STYLE, background="#ffffff", foreground="#718078", font=(FONT_NAME, 9))
        style.configure(SECTION_LABEL_STYLE, background="#ffffff", foreground="#19392f", font=(FONT_NAME, 12, "bold"))
        style.configure("TEntry", padding=(9, 8), fieldbackground="#f7f9f7")
        style.configure("TCombobox", padding=(8, 7), fieldbackground="#f7f9f7")
        style.configure("TButton", width=BUTTON_WIDTH, padding=(10, 7), anchor="center", font=(FONT_NAME, 9))
        style.configure(PRIMARY_BUTTON_STYLE, width=BUTTON_WIDTH, padding=(10, 7), anchor="center", background="#236b58", foreground="#ffffff", font=(FONT_NAME, 9, "bold"))
        style.map(PRIMARY_BUTTON_STYLE, background=[("active", "#1b5747")])
        style.configure(NEW_STORY_BUTTON_STYLE, width=BUTTON_WIDTH, padding=(10, 7), anchor="center", background="#2878c7", foreground="#ffffff", font=(FONT_NAME, 9, "bold"))
        style.map(NEW_STORY_BUTTON_STYLE, background=[("active", "#1f609f")])
        style.configure("Menu.TButton", width=24, padding=(10, 8), anchor="center", background="#236b58", foreground="#ffffff", font=(FONT_NAME, 9, "bold"))
        style.map("Menu.TButton", background=[("active", "#1b5747")])
        style.configure("Danger.TButton", width=BUTTON_WIDTH, padding=(10, 7), anchor="center", background="#c75050", foreground="#ffffff", font=(FONT_NAME, 9, "bold"))
        style.map("Danger.TButton", background=[("active", "#a53d3d")])
        style.configure("Header.TButton", width=BUTTON_WIDTH, padding=(10, 7), anchor="center", background="#315b4f", foreground="#ffffff", font=(FONT_NAME, 9, "bold"))
        style.map("Header.TButton", background=[("active", "#426f61")])
        style.configure(ICON_BUTTON_STYLE, width=ICON_BUTTON_WIDTH, padding=(8, 7), anchor="center", font=(FONT_NAME, 9))
        style.configure("Treeview", background="#ffffff", fieldbackground="#ffffff", foreground="#23352f", rowheight=34, font=(FONT_NAME, 9))
        style.configure("Treeview.Heading", background="#e7eeea", foreground="#365046", font=(FONT_NAME, 9, "bold"), padding=(8, 9))
        style.map("Treeview", background=[("selected", "#d9ebe3")], foreground=[("selected", "#19392f")])

    def _build_layout(self):
        self.header = tk.Frame(self.root, bg="#19392f", height=100)
        self.header.pack(fill="x")
        self.header.pack_propagate(False)
        title = tk.Frame(self.header, bg="#19392f")
        title.pack(side="left", padx=28, pady=18)
        tk.Label(title, text="SCRUM  /  PRODUCT BACKLOG", bg="#19392f", fg="#a9c9bb", font=(FONT_NAME, 9, "bold")).pack(anchor="w")
        tk.Label(title, text="Histórias do produto", bg="#19392f", fg="#ffffff", font=(FONT_NAME, 21, "bold")).pack(anchor="w", pady=(2, 0))
        self.count_label = tk.Label(self.header, text="0 ITENS", bg="#19392f", fg="#d7e8df", font=(FONT_NAME, 10, "bold"))
        self.count_label.pack(side="right", padx=30)
        ttk.Button(self.header, text="← Menu Principal", style="Header.TButton", command=self.show_main_menu).pack(side="right", padx=8)

        self.content = ttk.Frame(self.root, padding=(20, 18, 20, 20))
        self.content.pack(fill="both", expand=True)
        self.content.columnconfigure(0, weight=0, minsize=330)
        self.content.columnconfigure(1, weight=1)
        self.content.rowconfigure(0, weight=1)
        self.form_panel = ttk.Frame(self.content, style=PANEL_STYLE, padding=20)
        self.form_panel.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        self.form_panel.columnconfigure(0, weight=1)
        self.form_panel.rowconfigure(0, weight=1)
        form_canvas = tk.Canvas(self.form_panel, bg="#ffffff", highlightthickness=0)
        form_scrollbar = ttk.Scrollbar(self.form_panel, orient="vertical", command=form_canvas.yview)
        form_canvas.configure(yscrollcommand=form_scrollbar.set)
        form_canvas.grid(row=0, column=0, sticky="nsew")
        form_scrollbar.grid(row=0, column=1, sticky="ns")
        self.form_body = ttk.Frame(form_canvas, style=PANEL_STYLE)
        form_window = form_canvas.create_window((0, 0), window=self.form_body, anchor="nw")
        self.form_body.bind("<Configure>", lambda _event: form_canvas.configure(scrollregion=form_canvas.bbox("all")))
        form_canvas.bind("<Configure>", lambda event: form_canvas.itemconfigure(form_window, width=event.width))
        self.list_panel = ttk.Frame(self.content, style=PANEL_STYLE, padding=18)
        self.list_panel.grid(row=0, column=1, sticky="nsew", padx=(12, 0))
        self.list_panel.columnconfigure(0, weight=1)
        self.list_panel.rowconfigure(5, weight=1)
        self._build_form()
        self._build_list()

    def show_main_menu(self):
        if hasattr(self, "dashboard_form"):
            self.dashboard_form.frame.pack_forget()
        if hasattr(self, "sprint_form"):
            self.sprint_form.frame.pack_forget()
        if hasattr(self, "history_form"):
            self.history_form.frame.pack_forget()
        if hasattr(self, "governance_form"):
            self.governance_form.frame.pack_forget()
        if hasattr(self, "transparency_form"):
            self.transparency_form.frame.pack_forget()
        if hasattr(self, "quality_form"):
            self.quality_form.frame.pack_forget()
        self.header.pack_forget()
        self.content.pack_forget()
        if not hasattr(self, "menu_frame"):
            self.menu_frame = ttk.Frame(self.root, padding=24)
            self.menu_frame.pack(fill="both", expand=True)
            menu = ttk.Frame(self.menu_frame, style=PANEL_STYLE, padding=(30, 24))
            menu.pack(fill="both", expand=True)
            ttk.Label(menu, text="SCRUM  /  MENU PRINCIPAL", style=MUTED_LABEL_STYLE).pack(anchor="w")
            ttk.Label(menu, text="Formulários", style=SECTION_LABEL_STYLE).pack(anchor="w", pady=(4, 18))

            rows = ttk.Frame(menu, style=PANEL_STYLE)
            rows.pack(fill="both", expand=True)
            rows.columnconfigure(1, weight=1)
            form_actions = (
                ("Dashboards", self.show_dashboards),
                ("Product Backlog", self.show_backlog),
                ("Sprints Backlog", self.show_sprint_backlog),
                ("Histórico de Sprints", self.show_history_sprints),
                ("Governança", self.show_governance),
                ("Transparência", self.show_transparency),
                ("Qualidade", self.show_quality),
            )
            for row_index, (form_name, command) in enumerate(form_actions):
                ttk.Button(rows, text=form_name, style="Menu.TButton", command=command).grid(
                    row=row_index, column=0, sticky="w", padx=(0, 22), pady=5
                )
                ttk.Label(rows, text=FORM_SUMMARIES[form_name], wraplength=580).grid(
                    row=row_index, column=1, sticky="w", pady=5
                )

            ttk.Separator(menu).pack(fill="x", pady=(18, 12))
            manual_row = ttk.Frame(menu, style=PANEL_STYLE)
            manual_row.pack(fill="x")
            ttk.Label(manual_row, text="Manual de instruções").pack(side="left", padx=(0, 10))
            self.manual_var = tk.StringVar(value="Product Backlog")
            ttk.Combobox(manual_row, textvariable=self.manual_var, values=list(FORM_MANUALS), state="readonly", width=28).pack(side="left", padx=(0, 8))
            ttk.Button(manual_row, text="Abrir manual", style=PRIMARY_BUTTON_STYLE,
                       command=self.open_selected_manual).pack(side="left")
            footer = ttk.Frame(menu, style=PANEL_STYLE)
            footer.pack(fill="x", pady=(14, 0))
            ttk.Button(footer, text="Novo projeto", style="Danger.TButton",
                       command=self.start_new_project).pack(side="left")
            ttk.Button(footer, text="Sair", command=self.root.destroy).pack(side="right")
        else:
            self.menu_frame.pack(fill="both", expand=True)
        self.root.title("Scrum | Menu Principal")

    def open_selected_manual(self):
        form_name = self.manual_var.get()
        manual_text = FORM_MANUALS.get(form_name)
        if manual_text is None:
            messagebox.showwarning("Manual não encontrado", "Selecione um formulário da lista.", parent=self.root)
            return
        window = tk.Toplevel(self.root)
        window.title(f"Manual | {form_name}")
        window.geometry("720x560")
        window.transient(self.root)
        window.columnconfigure(0, weight=1)
        window.rowconfigure(1, weight=1)
        ttk.Label(window, text=form_name, style=SECTION_LABEL_STYLE).grid(row=0, column=0, sticky="w", padx=18, pady=(16, 8))
        text_frame = ttk.Frame(window)
        text_frame.grid(row=1, column=0, sticky="nsew", padx=18)
        text_frame.columnconfigure(0, weight=1)
        text_frame.rowconfigure(0, weight=1)
        manual = tk.Text(text_frame, wrap="word", relief="flat", padx=12, pady=10, bg="#f7f9f7", fg="#23352f", font=(FONT_NAME, 10))
        manual.insert("1.0", manual_text)
        manual.configure(state="disabled")
        manual.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(text_frame, orient="vertical", command=manual.yview)
        scroll.grid(row=0, column=1, sticky="ns")
        manual.configure(yscrollcommand=scroll.set)
        ttk.Button(window, text="Fechar", command=window.destroy).grid(row=2, column=0, sticky="e", padx=18, pady=14)

    def start_new_project(self):
        confirmed = messagebox.askyesno(
            "Iniciar novo projeto",
            "Apagar TODOS os dados do projeto atual?\n\n"
            "Serão removidos: histórias, tarefas, históricos, governança, transparência e qualidade.\n"
            "As tabelas de referência (Fibonacci e Status) serão mantidas.",
            icon="warning", parent=self.root,
        )
        if not confirmed:
            return
        try:
            reset_project_data()
        except sqlite3.Error as error:
            messagebox.showerror("Erro ao limpar dados", str(error), parent=self.root)
            return
        self.refresh_records()
        messagebox.showinfo("Novo projeto",
                            "Dados do projeto apagados. As referências foram mantidas.", parent=self.root)

    def show_sprint_backlog(self):
        self.menu_frame.pack_forget()
        if hasattr(self, "dashboard_form"):
            self.dashboard_form.frame.pack_forget()
        if hasattr(self, "history_form"):
            self.history_form.frame.pack_forget()
        if hasattr(self, "governance_form"):
            self.governance_form.frame.pack_forget()
        if hasattr(self, "transparency_form"):
            self.transparency_form.frame.pack_forget()
        if hasattr(self, "quality_form"):
            self.quality_form.frame.pack_forget()
        self.header.pack_forget()
        self.content.pack_forget()
        if not hasattr(self, "sprint_form"):
            from sprints_formulario import SprintBacklogForm

            self.sprint_form = SprintBacklogForm(self.root, self.show_main_menu)
        else:
            self.sprint_form.frame.pack(fill="both", expand=True)
            self.sprint_form.refresh_records()
        self.root.title("Scrum | Sprints Backlog")

    def show_history_sprints(self):
        self.menu_frame.pack_forget()
        self.header.pack_forget()
        self.content.pack_forget()
        if hasattr(self, "dashboard_form"):
            self.dashboard_form.frame.pack_forget()
        if hasattr(self, "sprint_form"):
            self.sprint_form.frame.pack_forget()
        if hasattr(self, "governance_form"):
            self.governance_form.frame.pack_forget()
        if hasattr(self, "transparency_form"):
            self.transparency_form.frame.pack_forget()
        if hasattr(self, "quality_form"):
            self.quality_form.frame.pack_forget()
        if not hasattr(self, "history_form"):
            from historico_formulario import HistoricoSprintsForm

            self.history_form = HistoricoSprintsForm(self.root, self.show_main_menu)
        else:
            self.history_form.frame.pack(fill="both", expand=True)
            self.history_form.refresh_records()
        self.root.title("Scrum | Histórico de Sprints")

    def show_governance(self):
        self.menu_frame.pack_forget()
        self.header.pack_forget()
        self.content.pack_forget()
        if hasattr(self, "dashboard_form"):
            self.dashboard_form.frame.pack_forget()
        if hasattr(self, "sprint_form"):
            self.sprint_form.frame.pack_forget()
        if hasattr(self, "history_form"):
            self.history_form.frame.pack_forget()
        if hasattr(self, "transparency_form"):
            self.transparency_form.frame.pack_forget()
        if hasattr(self, "quality_form"):
            self.quality_form.frame.pack_forget()
        if not hasattr(self, "governance_form"):
            from governanca_formulario import GovernancaForm

            self.governance_form = GovernancaForm(self.root, self.show_main_menu)
        else:
            self.governance_form.frame.pack(fill="both", expand=True)
            self.governance_form.refresh_records()
        self.root.title("Scrum | Governança")

    def show_transparency(self):
        self.menu_frame.pack_forget()
        self.header.pack_forget()
        self.content.pack_forget()
        for form_name in ("sprint_form", "history_form", "governance_form", "dashboard_form"):
            form = getattr(self, form_name, None)
            if form:
                form.frame.pack_forget()
        if not hasattr(self, "transparency_form"):
            from transparencia_formulario import TransparenciaForm

            self.transparency_form = TransparenciaForm(self.root, self.show_main_menu)
        else:
            self.transparency_form.frame.pack(fill="both", expand=True)
            self.transparency_form.refresh_records()
        self.root.title("Scrum | Transparência")

    def show_quality(self):
        self.menu_frame.pack_forget()
        self.header.pack_forget()
        self.content.pack_forget()
        for form_name in ("sprint_form", "history_form", "governance_form", "transparency_form", "dashboard_form"):
            form = getattr(self, form_name, None)
            if form:
                form.frame.pack_forget()
        if not hasattr(self, "quality_form"):
            from qualidade_formulario import QualidadeForm

            self.quality_form = QualidadeForm(self.root, self.show_main_menu)
        else:
            self.quality_form.frame.pack(fill="both", expand=True)
            self.quality_form.refresh_records()
        self.root.title("Scrum | Qualidade")

    def show_dashboards(self):
        self.menu_frame.pack_forget()
        self.header.pack_forget()
        self.content.pack_forget()
        for form_name in ("sprint_form", "history_form", "governance_form", "transparency_form", "quality_form"):
            form = getattr(self, form_name, None)
            if form:
                form.frame.pack_forget()
        if not hasattr(self, "dashboard_form"):
            from dashboard_formulario import DashboardForm

            self.dashboard_form = DashboardForm(self.root, self.show_main_menu)
        else:
            self.dashboard_form.frame.pack(fill="both", expand=True)
            self.dashboard_form.refresh_records()
        self.root.title("Scrum | Dashboards")

    def show_backlog(self):
        if hasattr(self, "dashboard_form"):
            self.dashboard_form.frame.pack_forget()
        if hasattr(self, "sprint_form"):
            self.sprint_form.frame.pack_forget()
        if hasattr(self, "history_form"):
            self.history_form.frame.pack_forget()
        if hasattr(self, "governance_form"):
            self.governance_form.frame.pack_forget()
        if hasattr(self, "transparency_form"):
            self.transparency_form.frame.pack_forget()
        if hasattr(self, "quality_form"):
            self.quality_form.frame.pack_forget()
        self.menu_frame.pack_forget()
        self.header.pack(fill="x")
        self.content.pack(fill="both", expand=True)
        self.root.title("Scrum | Product Backlog")

    def _build_form(self):
        ttk.Label(self.form_body, text="DETALHES DA HISTÓRIA", style=MUTED_LABEL_STYLE).pack(anchor="w")
        ttk.Label(self.form_body, text="Cadastro", style=SECTION_LABEL_STYLE).pack(anchor="w", pady=(3, 16))
        self.id_entry = self._add_entry("Id *", self.id_var, search_command=self.search_by_id)
        self.id_entry.configure(state="readonly")
        self.module_entry = self._add_entry(MODULE_LABEL, self.module_var, "Ex.: Usuário, Estoque", self.search_by_module)
        self._attach_uppercase(self.module_entry)

        story_heading = ttk.Frame(self.form_body, style=PANEL_STYLE)
        story_heading.pack(fill="x", pady=(11, 5))
        ttk.Label(story_heading, text="História de usuário").pack(side="left")
        ttk.Button(story_heading, text="ⓘ", style=ICON_BUTTON_STYLE, command=self.show_story_help).pack(side="right")
        self.story_text = self._add_text(4)
        self.story_text.bind("<KeyRelease>", self._uppercase_text)

        priority_row = ttk.Frame(self.form_body, style=PANEL_STYLE)
        priority_row.pack(fill="x", pady=(10, 0))
        ttk.Label(priority_row, text="Prioridade").pack(side="left")
        ttk.Button(priority_row, text="ⓘ", style=ICON_BUTTON_STYLE, command=self.show_priority_info).pack(side="right")
        self.priority_combo = ttk.Combobox(self.form_body, textvariable=self.priority_var,
                                            values=["ALTÍSSIMA", "ALTA", "MÉDIA", "BAIXA"], state="readonly")
        self.priority_combo.pack(fill="x", pady=(5, 0))
        self.priority_combo.bind("<<ComboboxSelected>>", self._priority_changed)

        self.points_combo = self._add_combo("Story Points", self.points_var, [])
        self._add_combo("Sprint alocada", self.sprint_var, self.sprint_options)
        self._add_status_picker()
        ttk.Label(self.form_body, text="Observações").pack(anchor="w", pady=(11, 5))
        self.notes_text = self._add_text(3)
        self.notes_text.bind("<KeyRelease>", self._uppercase_text)
        ttk.Separator(self.form_body).pack(fill="x", pady=18)
        ttk.Button(self.form_body, text="Nova história", style=NEW_STORY_BUTTON_STYLE,
               command=self.new_story).pack(fill="x", pady=(0, 8))
        self.save_button = ttk.Button(self.form_body, text="Salvar história", style="Primary.TButton", command=self.save_record)
        self.save_button.pack(fill="x")
        ttk.Button(self.form_body, text="Limpar formulário", command=self.clear_form).pack(fill="x", pady=(8, 0))
        ttk.Label(self.form_body, text=f"Id, {MODULE_LABEL.casefold()}, história e status são obrigatórios.", style=MUTED_LABEL_STYLE).pack(anchor="w", pady=(16, 0))

    def _add_entry(self, label, variable, placeholder=None, search_command=None):
        ttk.Label(self.form_body, text=label).pack(anchor="w", pady=(10, 5))
        field_row = ttk.Frame(self.form_body, style=PANEL_STYLE)
        field_row.pack(fill="x")
        entry = ttk.Entry(field_row, textvariable=variable)
        entry.pack(side="left", fill="x", expand=True)
        if search_command:
            ttk.Button(field_row, text="🔍", style=ICON_BUTTON_STYLE, command=search_command).pack(side="left", padx=(6, 0))
        if placeholder:
            self._set_placeholder(entry, variable, placeholder)
        return entry

    def _add_combo(self, label, variable, values):
        ttk.Label(self.form_body, text=label).pack(anchor="w", pady=(10, 5))
        combo = ttk.Combobox(self.form_body, textvariable=variable, values=values, state="readonly")
        combo.pack(fill="x")
        return combo

    def _add_status_picker(self):
        header = ttk.Frame(self.form_body, style=PANEL_STYLE)
        header.pack(fill="x", pady=(10, 5))
        ttk.Label(header, text="Status").pack(side="left")
        ttk.Button(header, text="ⓘ Orientação", command=self.show_status_info).pack(side="right")
        row = ttk.Frame(self.form_body, style=PANEL_STYLE)
        row.pack(fill="x")
        self.status_combo = ttk.Combobox(row, textvariable=self.status_var, values=self.status_options, state="readonly")
        self.status_combo.pack(side="left", fill="x", expand=True)

    @staticmethod
    def _set_placeholder(entry, variable, placeholder):
        entry.placeholder = placeholder
        entry.configure(style="Placeholder.TEntry")
        variable.set(placeholder)
        entry.bind("<FocusIn>", lambda _event: BacklogApp._clear_placeholder(entry, variable))
        entry.bind("<FocusOut>", lambda _event: BacklogApp._restore_placeholder(entry, variable))

    @staticmethod
    def _clear_placeholder(entry, variable):
        if variable.get() == entry.placeholder:
            variable.set("")
            entry.configure(style="TEntry")

    @staticmethod
    def _restore_placeholder(entry, variable):
        if not variable.get():
            variable.set(entry.placeholder)
            entry.configure(style="Placeholder.TEntry")

    def _attach_uppercase(self, entry):
        entry.bind("<KeyRelease>", lambda _event: self._uppercase_entry(entry))

    @staticmethod
    def _uppercase_entry(entry):
        value = entry.get()
        if value == getattr(entry, "placeholder", None):
            return
        uppercase_value = value.upper()
        if value != uppercase_value:
            cursor = entry.index(tk.INSERT)
            entry.delete(0, tk.END)
            entry.insert(0, uppercase_value)
            entry.icursor(min(cursor, len(uppercase_value)))

    @staticmethod
    def _uppercase_text(event):
        widget = event.widget
        value = widget.get("1.0", "end-1c")
        uppercase_value = value.upper()
        if value != uppercase_value:
            cursor = widget.index(tk.INSERT)
            widget.delete("1.0", "end")
            widget.insert("1.0", uppercase_value)
            widget.mark_set(tk.INSERT, cursor)

    def _add_text(self, height):
        widget = tk.Text(self.form_body, height=height, wrap="word", relief="flat", bd=0, padx=9, pady=8,
                 bg="#f7f9f7", fg="#23352f", insertbackground="#236b58", font=(FONT_NAME, 10),
                         highlightthickness=1, highlightbackground="#dce5df", highlightcolor="#6ca18c")
        widget.pack(fill="x")
        return widget

    def _build_list(self):
        ttk.Label(self.list_panel, text="VISÃO GERAL", style=MUTED_LABEL_STYLE).grid(row=0, column=0, sticky="w")
        ttk.Label(self.list_panel, text="Backlog", style=SECTION_LABEL_STYLE).grid(row=1, column=0, sticky="w", pady=(3, 12))
        filters = ttk.Frame(self.list_panel, style=PANEL_STYLE)
        filters.grid(row=2, column=0, sticky="ew")
        filters.columnconfigure(0, weight=1)
        search = ttk.Entry(filters, textvariable=self.search_var)
        search.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        search.bind("<KeyRelease>", self._on_grid_search)
        ttk.Combobox(filters, textvariable=self.status_filter_var, values=["TODOS", *self.status_options], state="readonly", width=18).grid(row=0, column=1, padx=(0, 8))
        ttk.Combobox(filters, textvariable=self.sprint_filter_var, values=["TODAS", *self.sprint_options], state="readonly", width=16).grid(row=0, column=2)
        self.status_filter_var.trace_add("write", lambda *_args: self.apply_filters())
        self.sprint_filter_var.trace_add("write", lambda *_args: self.apply_filters())

        toolbar = ttk.Frame(self.list_panel, style=PANEL_STYLE)
        toolbar.grid(row=3, column=0, sticky="ew", pady=(10, 10))
        ttk.Button(toolbar, text="Importar CSV", command=self.import_csv).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar, text="Importar Excel", command=self.import_excel).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar, text="Exportar CSV", command=self.export_csv).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar, text="Exportar Excel", command=self.export_excel).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar, text="Gerar modelos", command=self.generate_templates).pack(side="right")

        actions = ttk.Frame(self.list_panel, style=PANEL_STYLE)
        actions.grid(row=4, column=0, sticky="ew", pady=(0, 9))
        ttk.Button(actions, text="Editar selecionado", command=self.edit_selected).pack(side="left", padx=(0, 6))
        ttk.Button(actions, text="Excluir selecionado", command=self.delete_selected).pack(side="left", padx=(0, 6))

        table_frame = ttk.Frame(self.list_panel, style=PANEL_STYLE)
        table_frame.grid(row=5, column=0, sticky="nsew")
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        columns = ("id", "module", "story", "priority", "points", "sprint", "status")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")
        headings = {
            "id": ("Id", 90), "module": (MODULE_LABEL, 130), "story": ("História de usuário", 300),
            "priority": ("Prioridade", 110), "points": ("Pontos", 75), "sprint": ("Sprint", 110), "status": ("Status", 145),
        }
        for key, (label, width) in headings.items():
            self.tree.heading(key, text=label)
            self.tree.column(key, width=width, minwidth=60, anchor="w", stretch=key == "story")
        self.tree.tag_configure("even", background="#f7f9f7")
        self.tree.tag_configure("odd", background="#ffffff")
        self.tree.grid(row=0, column=0, sticky="nsew")
        self.empty_label = ttk.Label(table_frame, text="Nenhum cadastro encontrado para os filtros atuais.", style=MUTED_LABEL_STYLE)
        self.empty_label.place(relx=0.5, rely=0.5, anchor="center")
        vertical = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        horizontal.grid(row=1, column=0, sticky="ew")
        self.tree.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        self.tree.bind("<Double-1>", self.load_selected)
        ttk.Label(self.list_panel, text="Selecione uma linha ou dê dois cliques para editar.", style=MUTED_LABEL_STYLE).grid(row=6, column=0, sticky="w", pady=(9, 0))

    def refresh_records(self):
        try:
            self.records = list_product_backlog()
        except sqlite3.Error as error:
            messagebox.showerror("Erro ao carregar backlog", str(error), parent=self.root)
            self.records = []
        self.apply_filters()
        if not self.editing_id:
            self.id_var.set(self._next_id())
            self.id_entry.configure(state="readonly")

    def _next_id(self):
        largest = 0
        for record in self.records:
            match = re.fullmatch(r"US-(\d+)", str(record[0]).upper())
            if match:
                largest = max(largest, int(match.group(1)))
        return f"US-{largest + 1:02d}"

    def _on_grid_search(self, event=None):
        text = self.search_var.get()
        uppercase_text = text.upper()
        if text != uppercase_text:
            cursor = event.widget.index(tk.INSERT) if event else len(text)
            self.search_var.set(uppercase_text)
            if event:
                event.widget.icursor(min(cursor, len(uppercase_text)))
        self.search_mode = "all"
        self.apply_filters()

    def search_by_id(self):
        record_id = simpledialog.askstring("Pesquisar por Id", "Informe o Id (ex.: US-01):", parent=self.root)
        if record_id:
            self.search_var.set(record_id.strip())
            self.search_mode = "id"
            self.status_filter_var.set("TODOS")
            self.sprint_filter_var.set("TODAS")
            self.apply_filters()

    def search_by_module(self):
        module = self.module_var.get().strip()
        if module == getattr(self.module_entry, "placeholder", None):
            module = ""
        if not module:
            module = simpledialog.askstring("Pesquisar por módulo", "Informe o módulo:", parent=self.root) or ""
        if module:
            self.search_var.set(module.strip())
            self.search_mode = "module"
            self.status_filter_var.set("TODOS")
            self.sprint_filter_var.set("TODAS")
            self.apply_filters()

    def show_story_help(self):
        messagebox.showinfo(
            "História de usuário",
            "Estrutura sugerida:\nComo [perfil], quero [objetivo], para [benefício].\n\nExemplo: Como administrador, quero cadastrar produtos para controlar o estoque.",
            parent=self.root,
        )

    def show_priority_info(self):
        priority = self.priority_var.get()
        point_groups = {
            "ALTÍSSIMA": ("13", "21"),
            "ALTA": ("8",),
            "MÉDIA": ("3", "5"),
            "BAIXA": ("1", "2"),
        }
        selected_points = point_groups.get(priority, tuple(str(point) for point in self.points_options))
        details = [
            f"{points} pontos: {self.fibonacci_data.get(points, {}).get('Significado', 'Sem descrição cadastrada.') }\nExemplos: {self.fibonacci_data.get(points, {}).get('Exemplos', 'Sem exemplos cadastrados.') }"
            for points in selected_points
        ]
        title = f"Prioridade {priority}" if priority else "Escala Fibonacci"
        messagebox.showinfo(title, "\n\n".join(details), parent=self.root)

    def show_status_info(self):
        status = self.status_var.get().strip()
        if not status:
            messagebox.showinfo("Orientação de status", "Selecione um status para consultar o significado e quando usá-lo.", parent=self.root)
            return
        details = self.status_data.get(status)
        if details is None:
            details = next((value for name, value in self.status_data.items() if name.casefold() == status.casefold()), None)
        if details is None:
            messagebox.showwarning("Status não encontrado", f"Não há orientação cadastrada para '{status}'.", parent=self.root)
            return
        messagebox.showinfo(
            f"Status: {status}",
            f"Significado: {details.get('Significado') or 'Não informado.'}\n\nQuando usar: {details.get('Quando_usar') or 'Não informado.'}",
            parent=self.root,
        )

    def _priority_changed(self, event=None):
        point_groups = {
            "ALTÍSSIMA": ("13", "21"),
            "ALTA": ("8",),
            "MÉDIA": ("3", "5"),
            "BAIXA": ("1", "2"),
        }
        allowed = [point for point in point_groups.get(self.priority_var.get(), ()) if point in self.fibonacci_data]
        self.points_combo.configure(values=allowed)
        if event is not None and self.points_var.get() not in allowed:
            self.points_var.set("")

    def apply_filters(self):
        if not hasattr(self, "tree"):
            return
        query = self.search_var.get().strip().casefold()
        status = self.status_filter_var.get()
        sprint = self.sprint_filter_var.get()
        filtered = [record for record in self.records if self._matches_filters(record, query, status, sprint)]
        self.tree.delete(*self.tree.get_children())
        if filtered:
            self.empty_label.place_forget()
        else:
            self.empty_label.place(relx=0.5, rely=0.5, anchor="center")
        for index, record in enumerate(filtered):
            self.tree.insert("", "end", iid=str(index), values=(record[0], record[1], record[2], record[3], record[4], record[5], record[6]), tags=("even" if index % 2 == 0 else "odd",))
        grid_rows = [tuple(record[index] for index in range(7)) for record in self.records]
        fit_tree_columns(self.tree, self.root, BACKLOG_GRID_COLUMNS, grid_rows)
        self.count_label.configure(text=f"{len(filtered)} ITENS" if len(filtered) != 1 else "1 ITEM")

    def _matches_filters(self, record, query, status, sprint):
        if self.search_mode == "module":
            searchable = str(record[1] or "").casefold()
        elif self.search_mode == "id":
            searchable = str(record[0] or "").casefold()
        else:
            searchable = " ".join(str(value or "") for value in record[:3]).casefold()
        return (not query or query in searchable) and (status == "TODOS" or record[6] == status) and (sprint == "TODAS" or record[5] == sprint)

    def save_record(self):
        record_id = self.id_var.get().strip().upper()
        module = self.module_var.get().strip()
        if module == getattr(self.module_entry, "placeholder", None):
            module = ""
        story = self.story_text.get("1.0", "end-1c").strip()
        status = self.status_var.get().strip()
        missing = [label for label, value in (("Id", record_id), (MODULE_LABEL, module), ("História", story), ("Status", status)) if not value]
        if missing:
            messagebox.showwarning("Campos obrigatórios", f"Preencha: {', '.join(missing)}.", parent=self.root)
            return
        try:
            points = int(self.points_var.get()) if self.points_var.get() else None
        except ValueError:
            messagebox.showwarning("Story Points inválido", "Informe um número inteiro.", parent=self.root)
            return
        values = (module.upper(), story.upper(), self.priority_var.get().strip().upper(), points,
                  self.sprint_var.get().strip().upper(), status.upper(), self.notes_text.get("1.0", "end-1c").strip().upper())
        try:
            if self.editing_id:
                update_product_backlog(self.editing_id, *values)
                message = f"História {self.editing_id} atualizada."
            else:
                if record_id in {str(record[0]).upper() for record in self.records}:
                    messagebox.showwarning("Id já cadastrado", "Selecione a história existente para editá-la.", parent=self.root)
                    return
                insert_product_backlog(record_id, *values)
                message = f"História {record_id} cadastrada."
        except sqlite3.Error as error:
            messagebox.showerror("Não foi possível salvar", str(error), parent=self.root)
            return
        self.refresh_records()
        self.clear_form()
        messagebox.showinfo("Backlog atualizado", message, parent=self.root)

    def load_selected(self, _event=None):
        selection = self.tree.selection()
        if not selection:
            return
        selected_id = self.tree.item(selection[0], "values")[0]
        record = next((row for row in self.records if str(row[0]) == str(selected_id)), None)
        if record is None:
            return
        self.editing_id = record[0]
        self.id_var.set(record[0] or "")
        self.id_entry.configure(state="readonly")
        self.module_var.set(record[1] or "")
        self._set_text(self.story_text, record[2])
        self.priority_var.set(record[3] or "")
        self.points_var.set(str(record[4] or ""))
        self.sprint_var.set(record[5] or "")
        self.status_var.set(record[6] or "")
        self._set_text(self.notes_text, record[7])
        self._priority_changed()
        self.save_button.configure(text="Salvar alterações")

    def edit_selected(self):
        if not self.tree.selection():
            messagebox.showinfo("Selecione uma história", "Escolha uma linha do backlog para editar.", parent=self.root)
            return
        self.load_selected()

    def delete_selected(self):
        selection = self.tree.selection()
        record_id = self.tree.item(selection[0], "values")[0] if selection else self.editing_id
        if not record_id:
            messagebox.showinfo("Selecione uma história", "Escolha uma linha do backlog para excluir.", parent=self.root)
            return
        if not messagebox.askyesno("Confirmar exclusão", f"Excluir a história {record_id}?", parent=self.root):
            return
        try:
            delete_product_backlog(record_id)
        except sqlite3.Error as error:
            messagebox.showerror("Não foi possível excluir", str(error), parent=self.root)
            return
        self.refresh_records()
        self.clear_form()

    def new_story(self):
        self.clear_form(reset_filters=False)

    def clear_form(self, reset_filters=True):
        self.editing_id = None
        for variable in (self.module_var, self.priority_var, self.points_var, self.sprint_var, self.status_var):
            variable.set("")
        self._set_text(self.story_text, "")
        self._set_text(self.notes_text, "")
        self.module_var.set(self.module_entry.placeholder)
        self.module_entry.configure(style="Placeholder.TEntry")
        self.save_button.configure(text="Salvar história")
        self.tree.selection_remove(self.tree.selection())
        if reset_filters:
            self.search_mode = "all"
            self.search_var.set("")
            self.status_filter_var.set("TODOS")
            self.sprint_filter_var.set("TODAS")
        self.id_var.set(self._next_id())
        self.id_entry.configure(state="readonly")
        self.points_combo.configure(values=[])

    @staticmethod
    def _set_text(widget, value):
        widget.delete("1.0", "end")
        widget.insert("1.0", value or "")

    def _visible_records(self):
        visible_ids = {self.tree.item(item, "values")[0] for item in self.tree.get_children()}
        return [record for record in self.records if str(record[0]) in visible_ids]

    def export_csv(self):
        path = filedialog.asksaveasfilename(parent=self.root, title="Exportar backlog", defaultextension=".csv", filetypes=[("CSV", "*.csv")], initialfile="product_backlog.csv")
        if not path:
            return
        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as output:
                writer = csv.writer(output)
                writer.writerow(EXPORT_COLUMNS)
                writer.writerows(self._visible_records())
            messagebox.showinfo("Exportação concluída", f"Arquivo salvo em:\n{path}", parent=self.root)
        except OSError as error:
            messagebox.showerror("Erro ao exportar CSV", str(error), parent=self.root)

    def export_excel(self):
        path = filedialog.asksaveasfilename(parent=self.root, title="Exportar backlog", defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile="product_backlog.xlsx")
        if not path:
            return
        try:
            from openpyxl import Workbook
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Product Backlog"
            sheet.append(EXPORT_COLUMNS)
            for record in self._visible_records():
                sheet.append(list(record))
            workbook.save(path)
            messagebox.showinfo("Exportação concluída", f"Arquivo salvo em:\n{path}", parent=self.root)
        except ImportError:
            messagebox.showerror(DEPENDENCY_ERROR_TITLE, DEPENDENCY_ERROR_MESSAGE, parent=self.root)
        except (OSError, ValueError, zipfile.BadZipFile) as error:
            messagebox.showerror("Erro ao exportar Excel", str(error), parent=self.root)

    def import_csv(self):
        path = filedialog.askopenfilename(parent=self.root, title="Importar backlog", filetypes=[("CSV", "*.csv")])
        if not path:
            return
        try:
            with open(path, newline="", encoding="utf-8-sig") as source:
                self._import_rows(csv.DictReader(source))
        except (OSError, csv.Error, UnicodeError) as error:
            messagebox.showerror("Erro ao importar CSV", str(error), parent=self.root)

    def import_excel(self):
        path = filedialog.askopenfilename(parent=self.root, title="Importar backlog", filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            rows = workbook.active.iter_rows(values_only=True)
            headers = next(rows, ())
            self._import_rows(dict(zip(headers, row)) for row in rows)
            workbook.close()
        except ImportError:
            messagebox.showerror(DEPENDENCY_ERROR_TITLE, DEPENDENCY_ERROR_MESSAGE, parent=self.root)
        except (OSError, ValueError, zipfile.BadZipFile) as error:
            messagebox.showerror("Erro ao importar Excel", str(error), parent=self.root)

    def _import_rows(self, rows):
        aliases = {
            "id": ("id",), "module": (MODULE_LABEL.casefold(), "modulo"), "story": ("história", "historia", "historia_usuario"),
            "priority": ("prioridade",), "points": ("story points", "story_points"),
            "sprint": ("sprint", "sprint alocada", "sprint_alocada"), "status": ("status",),
            "notes": ("observações", "observacoes"),
        }
        existing = {str(record[0]).upper() for record in self.records}
        imported = skipped = 0
        for row in rows:
            normalized = {str(key).strip().casefold(): value for key, value in row.items() if key is not None}
            values = {}
            for field, names in aliases.items():
                value = next((normalized[name] for name in names if name in normalized), "")
                values[field] = "" if value is None else str(value).strip()
            if not values["id"] or not values["module"] or not values["story"] or not values["status"]:
                skipped += 1
                continue
            record_id = values["id"].upper()
            if record_id in existing:
                skipped += 1
                continue
            try:
                points = int(float(values["points"])) if values["points"] else None
                insert_product_backlog(record_id, values["module"].upper(), values["story"].upper(),
                                       values["priority"].upper(), points, values["sprint"].upper(),
                                       values["status"].upper(), values["notes"].upper())
                existing.add(record_id)
                imported += 1
            except (ValueError, TypeError, sqlite3.Error):
                skipped += 1
        self.refresh_records()
        messagebox.showinfo("Importação concluída", f"Importados: {imported}\nIgnorados ou inválidos: {skipped}", parent=self.root)

    def generate_templates(self):
        directory = filedialog.askdirectory(parent=self.root, title="Escolha onde salvar os modelos")
        if not directory:
            return
        csv_path = os.path.join(directory, "modelo_importacao_backlog.csv")
        excel_path = os.path.join(directory, "modelo_importacao_backlog.xlsx")
        try:
            with open(csv_path, "w", newline="", encoding="utf-8-sig") as output:
                csv.writer(output).writerow(EXPORT_COLUMNS)
            from openpyxl import Workbook
            workbook = Workbook()
            workbook.active.append(EXPORT_COLUMNS)
            workbook.save(excel_path)
            messagebox.showinfo("Modelos criados", f"{csv_path}\n{excel_path}", parent=self.root)
        except ImportError:
            messagebox.showerror(DEPENDENCY_ERROR_TITLE, DEPENDENCY_ERROR_MESSAGE, parent=self.root)
        except OSError as error:
            messagebox.showerror("Erro ao gerar modelos", str(error), parent=self.root)


def main(root):
    BacklogApp(root)
