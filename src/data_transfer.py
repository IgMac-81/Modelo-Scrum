import csv
import zipfile
from tkinter import filedialog, messagebox, ttk

DEPENDENCY_ERROR_TITLE = "Dependência ausente"
DEPENDENCY_ERROR_MESSAGE = "Instale as dependências com: pip install -r requirements.txt"


def add_transfer_buttons(parent, handler):
    for label, command in (
        ("Importar CSV", handler.import_csv),
        ("Importar Excel", handler.import_excel),
        ("Exportar CSV", handler.export_csv),
        ("Exportar Excel", handler.export_excel),
    ):
        ttk.Button(parent, text=label, command=command).pack(side="left", padx=(0, 6))


def export_csv(parent, title, initialfile, headers, rows):
    path = filedialog.asksaveasfilename(parent=parent, title=title, defaultextension=".csv",
                                        filetypes=[("CSV", "*.csv")], initialfile=initialfile)
    if not path:
        return
    try:
        with open(path, "w", newline="", encoding="utf-8-sig") as output:
            writer = csv.writer(output)
            writer.writerow(headers)
            writer.writerows(rows)
        messagebox.showinfo("Exportação concluída", f"Arquivo salvo em:\n{path}", parent=parent)
    except OSError as error:
        messagebox.showerror("Erro ao exportar CSV", str(error), parent=parent)


def export_excel(parent, title, initialfile, sheet_name, headers, rows):
    path = filedialog.asksaveasfilename(parent=parent, title=title, defaultextension=".xlsx",
                                        filetypes=[("Excel", "*.xlsx")], initialfile=initialfile)
    if not path:
        return
    try:
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(list(headers))
        for record in rows:
            sheet.append(["" if value is None else value for value in record])
        workbook.save(path)
        messagebox.showinfo("Exportação concluída", f"Arquivo salvo em:\n{path}", parent=parent)
    except ImportError:
        messagebox.showerror(DEPENDENCY_ERROR_TITLE, DEPENDENCY_ERROR_MESSAGE, parent=parent)
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        messagebox.showerror("Erro ao exportar Excel", str(error), parent=parent)


def read_csv(parent, title):
    path = filedialog.askopenfilename(parent=parent, title=title, filetypes=[("CSV", "*.csv")])
    if not path:
        return None
    try:
        with open(path, newline="", encoding="utf-8-sig") as source:
            return list(csv.DictReader(source))
    except (OSError, csv.Error, UnicodeError) as error:
        messagebox.showerror("Erro ao importar CSV", str(error), parent=parent)
        return None


def read_excel(parent, title):
    path = filedialog.askopenfilename(parent=parent, title=title, filetypes=[("Excel", "*.xlsx")])
    if not path:
        return None
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, ())
        data = [dict(zip(headers, row)) for row in rows]
        workbook.close()
        return data
    except ImportError:
        messagebox.showerror(DEPENDENCY_ERROR_TITLE, DEPENDENCY_ERROR_MESSAGE, parent=parent)
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        messagebox.showerror("Erro ao importar Excel", str(error), parent=parent)
    return None


def normalized(row):
    return {str(key).strip().casefold(): value for key, value in row.items() if key is not None}


def field(row, *names):
    value = next((row[name] for name in names if name in row), "")
    return "" if value is None else str(value).strip()


def report(parent, imported, skipped):
    messagebox.showinfo("Importação concluída",
                        f"Importados: {imported}\nIgnorados ou inválidos: {skipped}", parent=parent)
